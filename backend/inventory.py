"""
Smart Meter & Material Inventory Management — Session A
--------------------------------------------------------
Master Data + Smart Meter / Old Meter / Cable stock + Dashboard + Auto-movement ledger + Excel import.

Collections:
- inv_master (division/sub_division/sdo/store/agency/installer/meter_make/meter_model/cable_type/cable_size)
- inv_smart_meters
- inv_old_meters
- inv_cables (per drum / batch)
- inv_ledger (immutable audit / movement history)
"""
from fastapi import APIRouter, HTTPException, Depends, UploadFile, File, Query
from pydantic import BaseModel, Field, ConfigDict
from typing import Optional, List, Dict, Any, Literal
from datetime import datetime, timezone
import uuid
import io
import openpyxl
import logging

from auth import dep_current, dep_admin, _log_activity  # reuse existing auth

log = logging.getLogger(__name__)
_db = None

def init_inventory(db):
    global _db
    _db = db


inventory_router = APIRouter(prefix="/api/inventory", tags=["inventory"])

MASTER_TYPES = {"division", "sub_division", "sdo", "store", "agency",
                "installer", "meter_make", "meter_model", "cable_type",
                "cable_size", "document_type"}

SM_STATUSES = ("Available", "Issued", "Installed", "Returned", "Damaged", "Defective")
OM_CONDITIONS = ("Good", "Repairable", "Damaged", "Burnt", "Defective", "Scrap")
OM_DEPOSITS = ("Pending", "Deposited", "Verified")


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _clean(doc: dict) -> dict:
    if not doc:
        return doc
    doc.pop("_id", None)
    return doc


async def _ledger(actor_email: Optional[str], entity: str, entity_id: str,
                   action: str, before: Any = None, after: Any = None,
                   detail: str = ""):
    await _db.inv_ledger.insert_one({
        "id": str(uuid.uuid4()),
        "entity": entity,
        "entity_id": entity_id,
        "action": action,
        "before": before,
        "after": after,
        "detail": detail[:500],
        "actor": actor_email,
        "timestamp": _now(),
    })


# ================================================================
#                       MASTER DATA
# ================================================================
class MasterIn(BaseModel):
    type: Literal["division", "sub_division", "sdo", "store", "agency",
                  "installer", "meter_make", "meter_model", "cable_type",
                  "cable_size", "document_type"]
    name: str
    code: Optional[str] = None
    parent: Optional[str] = None  # e.g. sub_division has parent = division id
    active: bool = True
    extra: Optional[Dict[str, Any]] = None


@inventory_router.get("/masters")
async def list_masters(type: Optional[str] = None, user: dict = Depends(dep_current)):
    q = {}
    if type:
        if type not in MASTER_TYPES:
            raise HTTPException(400, "Unknown master type")
        q["type"] = type
    docs = await _db.inv_master.find(q, {"_id": 0}).sort([("type", 1), ("name", 1)]).to_list(2000)
    return {"items": docs}


@inventory_router.post("/masters")
async def add_master(payload: MasterIn, user: dict = Depends(dep_admin)):
    exists = await _db.inv_master.find_one({"type": payload.type,
                                             "name": {"$regex": f"^{payload.name.strip()}$",
                                                       "$options": "i"}})
    if exists:
        raise HTTPException(409, f"{payload.type} '{payload.name}' already exists")
    doc = {"id": str(uuid.uuid4()), **payload.model_dump(), "name": payload.name.strip(),
           "created_at": _now(), "created_by": user["email"]}
    await _db.inv_master.insert_one(doc)
    await _ledger(user["email"], "master", doc["id"], "created", None, doc)
    return _clean(doc)


@inventory_router.patch("/masters/{mid}")
async def update_master(mid: str, patch: Dict[str, Any], user: dict = Depends(dep_admin)):
    before = await _db.inv_master.find_one({"id": mid}, {"_id": 0})
    if not before:
        raise HTTPException(404, "Not found")
    allowed = {"name", "code", "parent", "active", "extra"}
    updates = {k: v for k, v in patch.items() if k in allowed}
    updates["updated_at"] = _now()
    await _db.inv_master.update_one({"id": mid}, {"$set": updates})
    after = await _db.inv_master.find_one({"id": mid}, {"_id": 0})
    await _ledger(user["email"], "master", mid, "updated", before, after)
    return after


@inventory_router.delete("/masters/{mid}")
async def delete_master(mid: str, user: dict = Depends(dep_admin)):
    doc = await _db.inv_master.find_one({"id": mid})
    if not doc:
        raise HTTPException(404, "Not found")
    # Soft delete = mark inactive
    await _db.inv_master.update_one({"id": mid}, {"$set": {"active": False, "deleted_at": _now()}})
    await _ledger(user["email"], "master", mid, "deleted", doc, None)
    return {"ok": True}


# ================================================================
#                    SMART METER INVENTORY
# ================================================================
SM_FIELDS = ["serial_number", "meter_make", "meter_model", "meter_type",
             "rating", "batch_number", "purchase_date", "received_qty",
             "status", "division", "sub_division", "sdo", "store_location",
             "vendor", "remarks"]


class SmartMeterIn(BaseModel):
    serial_number: str
    meter_make: Optional[str] = None
    meter_model: Optional[str] = None
    meter_type: Optional[str] = None
    rating: Optional[str] = None
    batch_number: Optional[str] = None
    purchase_date: Optional[str] = None
    received_qty: int = 1
    status: Literal["Available", "Issued", "Installed", "Returned", "Damaged", "Defective"] = "Available"
    division: Optional[str] = None
    sub_division: Optional[str] = None
    sdo: Optional[str] = None
    store_location: Optional[str] = None
    vendor: Optional[str] = None
    remarks: Optional[str] = None


@inventory_router.get("/smart-meters")
async def list_smart_meters(
    q: Optional[str] = None,
    status: Optional[str] = None,
    division: Optional[str] = None,
    make: Optional[str] = None,
    page: int = 1, page_size: int = Query(25, le=200),
    user: dict = Depends(dep_current),
):
    query: Dict[str, Any] = {"deleted_at": {"$exists": False}}
    if status: query["status"] = status
    if division: query["division"] = division
    if make: query["meter_make"] = make
    if q:
        rex = {"$regex": q, "$options": "i"}
        query["$or"] = [{"serial_number": rex}, {"batch_number": rex},
                        {"meter_make": rex}, {"meter_model": rex}, {"vendor": rex}]
    skip = (max(1, page) - 1) * page_size
    total = await _db.inv_smart_meters.count_documents(query)
    docs = await _db.inv_smart_meters.find(query, {"_id": 0}).sort("created_at", -1).skip(skip).limit(page_size).to_list(page_size)
    return {"items": docs, "total": total, "page": page, "page_size": page_size}


@inventory_router.post("/smart-meters")
async def add_smart_meter(payload: SmartMeterIn, user: dict = Depends(dep_admin)):
    sn = payload.serial_number.strip().upper()
    if not sn:
        raise HTTPException(400, "Serial number required")
    if await _db.inv_smart_meters.find_one({"serial_number": sn, "deleted_at": {"$exists": False}}):
        raise HTTPException(409, f"Duplicate serial number: {sn}")
    doc = {"id": str(uuid.uuid4()), **payload.model_dump(),
           "serial_number": sn, "created_at": _now(), "created_by": user["email"]}
    await _db.inv_smart_meters.insert_one(doc)
    await _ledger(user["email"], "smart_meter", doc["id"], "created", None,
                   {"serial": sn, "status": doc["status"]}, "Received")
    return _clean(doc)


@inventory_router.get("/smart-meters/{mid}")
async def get_smart_meter(mid: str, user: dict = Depends(dep_current)):
    doc = await _db.inv_smart_meters.find_one({"id": mid}, {"_id": 0})
    if not doc:
        raise HTTPException(404, "Not found")
    # Attach history — all ledger entries for this meter
    history = await _db.inv_ledger.find({"entity": "smart_meter", "entity_id": mid},
                                         {"_id": 0}).sort("timestamp", 1).to_list(200)
    doc["history"] = history
    return doc


@inventory_router.patch("/smart-meters/{mid}")
async def update_smart_meter(mid: str, patch: Dict[str, Any], user: dict = Depends(dep_admin)):
    before = await _db.inv_smart_meters.find_one({"id": mid}, {"_id": 0})
    if not before:
        raise HTTPException(404, "Not found")
    allowed = set(SM_FIELDS)
    updates = {k: v for k, v in patch.items() if k in allowed}
    if "serial_number" in updates:
        sn = updates["serial_number"].strip().upper()
        if sn != before["serial_number"]:
            if await _db.inv_smart_meters.find_one({"serial_number": sn, "id": {"$ne": mid}}):
                raise HTTPException(409, f"Duplicate serial number: {sn}")
            updates["serial_number"] = sn
    updates["updated_at"] = _now()
    updates["updated_by"] = user["email"]
    await _db.inv_smart_meters.update_one({"id": mid}, {"$set": updates})
    after = await _db.inv_smart_meters.find_one({"id": mid}, {"_id": 0})
    await _ledger(user["email"], "smart_meter", mid, "updated", before, after)
    return after


@inventory_router.delete("/smart-meters/{mid}")
async def delete_smart_meter(mid: str, user: dict = Depends(dep_admin)):
    doc = await _db.inv_smart_meters.find_one({"id": mid})
    if not doc:
        raise HTTPException(404, "Not found")
    await _db.inv_smart_meters.update_one({"id": mid}, {"$set": {"deleted_at": _now()}})
    await _ledger(user["email"], "smart_meter", mid, "deleted", doc, None)
    return {"ok": True}


# ================================================================
#                    OLD DIGITAL METER INVENTORY
# ================================================================
class OldMeterIn(BaseModel):
    serial_number: str
    meter_make: Optional[str] = None
    meter_model: Optional[str] = None
    meter_type: Optional[str] = None
    consumer_number: Optional[str] = None
    consumer_name: Optional[str] = None
    removal_date: Optional[str] = None
    division: Optional[str] = None
    sub_division: Optional[str] = None
    sdo: Optional[str] = None
    installer: Optional[str] = None
    condition: Literal["Good", "Repairable", "Damaged", "Burnt", "Defective", "Scrap"] = "Good"
    deposit_status: Literal["Pending", "Deposited", "Verified"] = "Pending"
    deposit_date: Optional[str] = None
    store_location: Optional[str] = None
    remarks: Optional[str] = None


@inventory_router.get("/old-meters")
async def list_old_meters(
    q: Optional[str] = None,
    deposit_status: Optional[str] = None,
    condition: Optional[str] = None,
    division: Optional[str] = None,
    page: int = 1, page_size: int = Query(25, le=200),
    user: dict = Depends(dep_current),
):
    query: Dict[str, Any] = {"deleted_at": {"$exists": False}}
    if deposit_status: query["deposit_status"] = deposit_status
    if condition: query["condition"] = condition
    if division: query["division"] = division
    if q:
        rex = {"$regex": q, "$options": "i"}
        query["$or"] = [{"serial_number": rex}, {"consumer_number": rex},
                        {"consumer_name": rex}, {"meter_make": rex}]
    skip = (max(1, page) - 1) * page_size
    total = await _db.inv_old_meters.count_documents(query)
    docs = await _db.inv_old_meters.find(query, {"_id": 0}).sort("removal_date", -1).skip(skip).limit(page_size).to_list(page_size)
    return {"items": docs, "total": total, "page": page, "page_size": page_size}


@inventory_router.post("/old-meters")
async def add_old_meter(payload: OldMeterIn, user: dict = Depends(dep_admin)):
    sn = payload.serial_number.strip().upper()
    if await _db.inv_old_meters.find_one({"serial_number": sn, "deleted_at": {"$exists": False}}):
        raise HTTPException(409, f"Duplicate old meter serial: {sn}")
    doc = {"id": str(uuid.uuid4()), **payload.model_dump(),
           "serial_number": sn, "created_at": _now(), "created_by": user["email"]}
    await _db.inv_old_meters.insert_one(doc)
    await _ledger(user["email"], "old_meter", doc["id"], "created", None, doc, "Removed & received in store")
    return _clean(doc)


@inventory_router.patch("/old-meters/{mid}")
async def update_old_meter(mid: str, patch: Dict[str, Any], user: dict = Depends(dep_admin)):
    before = await _db.inv_old_meters.find_one({"id": mid}, {"_id": 0})
    if not before:
        raise HTTPException(404, "Not found")
    updates = {k: v for k, v in patch.items() if k not in ("id", "_id", "created_at", "created_by")}
    updates["updated_at"] = _now()
    updates["updated_by"] = user["email"]
    # Auto-set deposit_date when moving to Deposited
    if updates.get("deposit_status") == "Deposited" and not before.get("deposit_date"):
        updates["deposit_date"] = updates.get("deposit_date") or datetime.now(timezone.utc).date().isoformat()
    await _db.inv_old_meters.update_one({"id": mid}, {"$set": updates})
    after = await _db.inv_old_meters.find_one({"id": mid}, {"_id": 0})
    await _ledger(user["email"], "old_meter", mid, "updated", before, after)
    return after


@inventory_router.delete("/old-meters/{mid}")
async def delete_old_meter(mid: str, user: dict = Depends(dep_admin)):
    doc = await _db.inv_old_meters.find_one({"id": mid})
    if not doc:
        raise HTTPException(404, "Not found")
    await _db.inv_old_meters.update_one({"id": mid}, {"$set": {"deleted_at": _now()}})
    await _ledger(user["email"], "old_meter", mid, "deleted", doc, None)
    return {"ok": True}


# ================================================================
#                    CABLE INVENTORY (per drum)
# ================================================================
class CableIn(BaseModel):
    cable_type: str
    cable_size: str
    cable_spec: Optional[str] = None
    make: Optional[str] = None
    batch_number: Optional[str] = None
    drum_number: str
    unit: str = "meter"
    opening_stock: float = 0
    received_qty: float = 0
    issued_qty: float = 0
    used_qty: float = 0
    returned_qty: float = 0
    damaged_qty: float = 0
    division: Optional[str] = None
    sub_division: Optional[str] = None
    sdo: Optional[str] = None
    store_location: Optional[str] = None
    receipt_date: Optional[str] = None
    remarks: Optional[str] = None


def _cable_balance(d: dict) -> float:
    return round(
        (d.get("opening_stock") or 0)
        + (d.get("received_qty") or 0)
        + (d.get("returned_qty") or 0)
        - (d.get("issued_qty") or 0)
        - (d.get("damaged_qty") or 0),
        3,
    )


@inventory_router.get("/cables")
async def list_cables(
    q: Optional[str] = None,
    cable_type: Optional[str] = None,
    division: Optional[str] = None,
    page: int = 1, page_size: int = Query(25, le=200),
    user: dict = Depends(dep_current),
):
    query: Dict[str, Any] = {"deleted_at": {"$exists": False}}
    if cable_type: query["cable_type"] = cable_type
    if division: query["division"] = division
    if q:
        rex = {"$regex": q, "$options": "i"}
        query["$or"] = [{"drum_number": rex}, {"batch_number": rex},
                        {"cable_type": rex}, {"cable_size": rex}, {"make": rex}]
    skip = (max(1, page) - 1) * page_size
    total = await _db.inv_cables.count_documents(query)
    docs = await _db.inv_cables.find(query, {"_id": 0}).sort("created_at", -1).skip(skip).limit(page_size).to_list(page_size)
    for d in docs:
        d["balance_qty"] = _cable_balance(d)
    return {"items": docs, "total": total, "page": page, "page_size": page_size}


@inventory_router.post("/cables")
async def add_cable(payload: CableIn, user: dict = Depends(dep_admin)):
    drum = payload.drum_number.strip().upper()
    if not drum:
        raise HTTPException(400, "Drum number required")
    if await _db.inv_cables.find_one({"drum_number": drum, "deleted_at": {"$exists": False}}):
        raise HTTPException(409, f"Duplicate drum: {drum}")
    doc = {"id": str(uuid.uuid4()), **payload.model_dump(),
           "drum_number": drum, "created_at": _now(), "created_by": user["email"]}
    doc["balance_qty"] = _cable_balance(doc)
    await _db.inv_cables.insert_one(doc)
    await _ledger(user["email"], "cable", doc["id"], "created", None,
                   {"drum": drum, "opening": doc["opening_stock"], "received": doc["received_qty"]})
    return _clean(doc)


@inventory_router.patch("/cables/{cid}")
async def update_cable(cid: str, patch: Dict[str, Any], user: dict = Depends(dep_admin)):
    before = await _db.inv_cables.find_one({"id": cid}, {"_id": 0})
    if not before:
        raise HTTPException(404, "Not found")
    updates = {k: v for k, v in patch.items() if k not in ("id", "_id", "created_at", "created_by")}
    updates["updated_at"] = _now()
    updates["updated_by"] = user["email"]
    await _db.inv_cables.update_one({"id": cid}, {"$set": updates})
    after = await _db.inv_cables.find_one({"id": cid}, {"_id": 0})
    after["balance_qty"] = _cable_balance(after)
    await _db.inv_cables.update_one({"id": cid}, {"$set": {"balance_qty": after["balance_qty"]}})
    await _ledger(user["email"], "cable", cid, "updated", before, after)
    return after


@inventory_router.delete("/cables/{cid}")
async def delete_cable(cid: str, user: dict = Depends(dep_admin)):
    doc = await _db.inv_cables.find_one({"id": cid})
    if not doc:
        raise HTTPException(404, "Not found")
    await _db.inv_cables.update_one({"id": cid}, {"$set": {"deleted_at": _now()}})
    await _ledger(user["email"], "cable", cid, "deleted", doc, None)
    return {"ok": True}


# ================================================================
#                     DASHBOARD (aggregated KPIs)
# ================================================================
@inventory_router.get("/dashboard")
async def dashboard(user: dict = Depends(dep_current)):
    async def _count_sm(status=None):
        q = {"deleted_at": {"$exists": False}}
        if status: q["status"] = status
        return await _db.inv_smart_meters.count_documents(q)

    async def _count_om(**kw):
        q = {"deleted_at": {"$exists": False}, **kw}
        return await _db.inv_old_meters.count_documents(q)

    sm_total = await _count_sm()
    smart = {
        "total": sm_total,
        "available": await _count_sm("Available"),
        "issued": await _count_sm("Issued"),
        "installed": await _count_sm("Installed"),
        "returned": await _count_sm("Returned"),
        "damaged": await _count_sm("Damaged"),
    }
    old = {
        "total_removed": await _count_om(),
        "deposited": await _count_om(deposit_status="Deposited"),
        "pending_deposit": await _count_om(deposit_status="Pending"),
        "damaged": await _count_om(condition="Damaged"),
        "available_in_store": await _count_om(deposit_status="Pending"),
    }

    # Cable aggregates
    cable_agg = await _db.inv_cables.aggregate([
        {"$match": {"deleted_at": {"$exists": False}}},
        {"$group": {"_id": None,
                    "received": {"$sum": "$received_qty"},
                    "issued": {"$sum": "$issued_qty"},
                    "used": {"$sum": "$used_qty"},
                    "returned": {"$sum": "$returned_qty"},
                    "damaged": {"$sum": "$damaged_qty"},
                    "opening": {"$sum": "$opening_stock"}}}
    ]).to_list(1)
    ca = cable_agg[0] if cable_agg else {"received": 0, "issued": 0, "used": 0,
                                          "returned": 0, "damaged": 0, "opening": 0}
    balance = (ca["opening"] + ca["received"] + ca["returned"]
               - ca["issued"] - ca["damaged"])
    cable = {
        "total_received": round(ca["received"] + ca["opening"], 2),
        "total_issued": round(ca["issued"], 2),
        "total_used": round(ca["used"], 2),
        "available_balance": round(balance, 2),
        "returned": round(ca["returned"], 2),
        "damaged": round(ca["damaged"], 2),
    }

    # Division-wise smart meter stock
    div_agg = await _db.inv_smart_meters.aggregate([
        {"$match": {"deleted_at": {"$exists": False}}},
        {"$group": {"_id": "$division", "count": {"$sum": 1}}},
        {"$sort": {"count": -1}},
        {"$limit": 8},
    ]).to_list(20)
    division_stock = [{"division": d["_id"] or "—", "count": d["count"]} for d in div_agg]

    # Status distribution for pie
    status_dist = [
        {"name": "Available", "value": smart["available"]},
        {"name": "Issued", "value": smart["issued"]},
        {"name": "Installed", "value": smart["installed"]},
        {"name": "Returned", "value": smart["returned"]},
        {"name": "Damaged", "value": smart["damaged"]},
    ]

    # Recent ledger transactions
    recent = await _db.inv_ledger.find({}, {"_id": 0}).sort("timestamp", -1).to_list(15)

    return {
        "smart": smart, "old": old, "cable": cable,
        "documents": {"total_gate_pass": 0, "gate_pass_pending": 0, "gate_pass_verified": 0,
                       "total_bisignoff": 0, "bisignoff_pending": 0, "bisignoff_verified": 0},
        "division_stock": division_stock,
        "status_distribution": status_dist,
        "recent_transactions": recent,
    }


# ================================================================
#                     EXCEL IMPORT (per collection)
# ================================================================
@inventory_router.post("/import/smart-meters")
async def import_smart_meters(file: UploadFile = File(...), user: dict = Depends(dep_admin)):
    return await _generic_import(file, "inv_smart_meters", "smart_meter",
                                  key_field="serial_number", user=user,
                                  required=["serial_number"],
                                  transform_key=lambda v: str(v).strip().upper())


@inventory_router.post("/import/old-meters")
async def import_old_meters(file: UploadFile = File(...), user: dict = Depends(dep_admin)):
    return await _generic_import(file, "inv_old_meters", "old_meter",
                                  key_field="serial_number", user=user,
                                  required=["serial_number"],
                                  transform_key=lambda v: str(v).strip().upper())


@inventory_router.post("/import/cables")
async def import_cables(file: UploadFile = File(...), user: dict = Depends(dep_admin)):
    return await _generic_import(file, "inv_cables", "cable",
                                  key_field="drum_number", user=user,
                                  required=["drum_number", "cable_type", "cable_size"],
                                  transform_key=lambda v: str(v).strip().upper())


@inventory_router.post("/import/masters")
async def import_masters(type: str, file: UploadFile = File(...), user: dict = Depends(dep_admin)):
    if type not in MASTER_TYPES:
        raise HTTPException(400, "Unknown master type")
    content = await file.read()
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        raise HTTPException(400, "Empty file")
    headers = [str(h).strip().lower().replace(" ", "_") if h else "" for h in rows[0]]
    imported = 0; skipped = 0; errors = []
    for i, r in enumerate(rows[1:], start=2):
        rec = {h: (str(v).strip() if v is not None else "") for h, v in zip(headers, r) if h}
        name = rec.get("name") or rec.get(type) or rec.get(type.replace("_", ""))
        if not name:
            skipped += 1; continue
        exists = await _db.inv_master.find_one({"type": type,
                                                 "name": {"$regex": f"^{name}$", "$options": "i"}})
        if exists:
            skipped += 1; continue
        doc = {"id": str(uuid.uuid4()), "type": type, "name": name,
               "code": rec.get("code"), "active": True,
               "created_at": _now(), "created_by": user["email"]}
        try:
            await _db.inv_master.insert_one(doc)
            imported += 1
        except Exception as e:
            errors.append({"row": i, "error": str(e)[:200]})
    return {"imported": imported, "skipped": skipped, "errors": errors, "type": type}


async def _generic_import(file: UploadFile, coll: str, entity: str,
                           key_field: str, user: dict, required: List[str],
                           transform_key=lambda v: v):
    content = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    except Exception:
        raise HTTPException(400, "Cannot read Excel file")
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows or len(rows) < 2:
        raise HTTPException(400, "Empty file")
    headers = [str(h).strip().lower().replace(" ", "_").replace("-", "_") if h else "" for h in rows[0]]

    imported = 0; skipped = 0; errors = []
    coll_ref = getattr(_db, coll)
    for i, r in enumerate(rows[1:], start=2):
        rec = {h: v for h, v in zip(headers, r) if h}
        # Normalise empties
        rec = {k: (str(v).strip() if isinstance(v, str) else v) for k, v in rec.items()
                if v is not None and v != ""}
        # Required fields
        missing = [f for f in required if not rec.get(f)]
        if missing:
            errors.append({"row": i, "error": f"Missing {missing}"}); skipped += 1; continue
        key = transform_key(rec[key_field])
        rec[key_field] = key
        # Dedupe
        if await coll_ref.find_one({key_field: key, "deleted_at": {"$exists": False}}):
            skipped += 1; continue
        doc = {"id": str(uuid.uuid4()), **rec,
               "created_at": _now(), "created_by": user["email"], "import": True}
        if entity == "cable":
            for numf in ("opening_stock", "received_qty", "issued_qty", "used_qty",
                          "returned_qty", "damaged_qty"):
                try: doc[numf] = float(doc.get(numf) or 0)
                except Exception: doc[numf] = 0
            doc["balance_qty"] = _cable_balance(doc)
        try:
            await coll_ref.insert_one(doc)
            await _ledger(user["email"], entity, doc["id"], "imported", None,
                           {"key": key}, f"Row {i}")
            imported += 1
        except Exception as e:
            errors.append({"row": i, "error": str(e)[:200]})
    return {"imported": imported, "skipped": skipped, "errors": errors,
            "sample_columns": headers[:15]}


# ================================================================
#                     LEDGER (view)
# ================================================================
@inventory_router.get("/ledger")
async def get_ledger(
    entity: Optional[str] = None,
    entity_id: Optional[str] = None,
    limit: int = Query(100, le=500),
    user: dict = Depends(dep_current),
):
    q = {}
    if entity: q["entity"] = entity
    if entity_id: q["entity_id"] = entity_id
    docs = await _db.inv_ledger.find(q, {"_id": 0}).sort("timestamp", -1).to_list(limit)
    return {"items": docs}


# ================================================================
#            SESSION B — TRANSACTIONS (lifecycle)
# ================================================================
# 4 new collections: installations, gate_passes, cable_issues, bisignoffs
# Each mutation runs "auto-movement" that updates stock status.
# Documents (photos / PDFs) are stored as base64 data URLs on the doc.


class InstallationIn(BaseModel):
    installation_date: str
    division: Optional[str] = None
    sub_division: Optional[str] = None
    sdo: Optional[str] = None
    consumer_name: str
    consumer_number: str
    address: Optional[str] = None
    old_meter_serial: Optional[str] = None
    new_meter_serial: str
    installer: Optional[str] = None
    agency: Optional[str] = None
    mobile_number: Optional[str] = None
    status: Literal["Pending", "Installed", "Rejected", "Revisit Required", "Completed"] = "Installed"
    latitude: Optional[float] = None
    longitude: Optional[float] = None
    before_photo: Optional[str] = None       # base64 data URL
    after_photo: Optional[str] = None
    meter_photo: Optional[str] = None
    remarks: Optional[str] = None


@inventory_router.post("/installations")
async def add_installation(payload: InstallationIn, user: dict = Depends(dep_admin)):
    new_sn = payload.new_meter_serial.strip().upper()
    meter = await _db.inv_smart_meters.find_one({"serial_number": new_sn,
                                                   "deleted_at": {"$exists": False}})
    if not meter:
        raise HTTPException(400, f"Smart meter {new_sn} not in stock")
    if meter["status"] in ("Installed", "Damaged", "Returned"):
        raise HTTPException(400, f"Meter {new_sn} is {meter['status']} — not installable")

    doc = {"id": str(uuid.uuid4()), **payload.model_dump(),
           "new_meter_serial": new_sn,
           "old_meter_serial": (payload.old_meter_serial or "").strip().upper() or None,
           "created_at": _now(), "created_by": user["email"]}
    await _db.inv_installations.insert_one(doc)

    # Auto-movement: set smart meter to Installed
    await _db.inv_smart_meters.update_one({"id": meter["id"]},
        {"$set": {"status": "Installed", "installed_at": _now(),
                  "installation_id": doc["id"]}})
    await _ledger(user["email"], "smart_meter", meter["id"], "installed",
                   {"status": meter["status"]}, {"status": "Installed"},
                   f"Installed at consumer {payload.consumer_number} ({payload.consumer_name})")

    # Auto-movement: if old serial provided, add it to old-meter stock (Pending deposit)
    if doc["old_meter_serial"]:
        exists = await _db.inv_old_meters.find_one({"serial_number": doc["old_meter_serial"]})
        if not exists:
            om = {"id": str(uuid.uuid4()), "serial_number": doc["old_meter_serial"],
                   "consumer_number": payload.consumer_number,
                   "consumer_name": payload.consumer_name,
                   "removal_date": payload.installation_date,
                   "division": payload.division, "sub_division": payload.sub_division,
                   "sdo": payload.sdo, "installer": payload.installer or payload.agency,
                   "condition": "Good", "deposit_status": "Pending",
                   "installation_id": doc["id"],
                   "created_at": _now(), "created_by": user["email"]}
            await _db.inv_old_meters.insert_one(om)
            await _ledger(user["email"], "old_meter", om["id"], "removed",
                           None, {"serial": om["serial_number"]},
                           f"Auto-removed via installation {doc['id']}")

    await _ledger(user["email"], "installation", doc["id"], "created", None,
                   {"consumer": payload.consumer_number, "new_serial": new_sn})
    return _clean(doc)


@inventory_router.get("/installations")
async def list_installations(
    q: Optional[str] = None, status: Optional[str] = None,
    division: Optional[str] = None,
    page: int = 1, page_size: int = Query(25, le=200),
    user: dict = Depends(dep_current),
):
    query = {"deleted_at": {"$exists": False}}
    if status: query["status"] = status
    if division: query["division"] = division
    if q:
        rex = {"$regex": q, "$options": "i"}
        query["$or"] = [{"consumer_number": rex}, {"consumer_name": rex},
                        {"new_meter_serial": rex}, {"old_meter_serial": rex},
                        {"installer": rex}, {"mobile_number": rex}]
    skip = (max(1, page) - 1) * page_size
    total = await _db.inv_installations.count_documents(query)
    docs = await _db.inv_installations.find(query, {"_id": 0}).sort("installation_date", -1).skip(skip).limit(page_size).to_list(page_size)
    return {"items": docs, "total": total, "page": page, "page_size": page_size}


@inventory_router.get("/installations/{iid}")
async def get_installation(iid: str, user: dict = Depends(dep_current)):
    doc = await _db.inv_installations.find_one({"id": iid}, {"_id": 0})
    if not doc: raise HTTPException(404, "Not found")
    return doc


# ---------- GATE PASS ----------
class GatePassIn(BaseModel):
    gate_pass_number: str
    gate_pass_date: str
    division: Optional[str] = None
    sub_division: Optional[str] = None
    sdo: Optional[str] = None
    from_location: str
    to_location: str
    meter_serials: List[str] = []
    vehicle_number: Optional[str] = None
    driver_name: Optional[str] = None
    driver_mobile: Optional[str] = None
    agency: Optional[str] = None
    purpose: Optional[str] = None
    prepared_by: Optional[str] = None
    approved_by: Optional[str] = None
    status: Literal["Draft", "Submitted", "Approved", "Rejected", "Completed"] = "Draft"
    document: Optional[str] = None            # base64 PDF or image
    document_name: Optional[str] = None
    remarks: Optional[str] = None


@inventory_router.post("/gate-passes")
async def add_gate_pass(payload: GatePassIn, user: dict = Depends(dep_admin)):
    if not payload.gate_pass_number.strip():
        raise HTTPException(400, "Gate pass number required")
    if await _db.inv_gate_passes.find_one({"gate_pass_number": payload.gate_pass_number.strip(),
                                             "deleted_at": {"$exists": False}}):
        raise HTTPException(409, "Duplicate gate pass number")

    # Serial number validation — must exist and NOT be Issued/Installed/Damaged/Returned
    serials = [s.strip().upper() for s in payload.meter_serials if s and s.strip()]
    invalid = []
    for sn in serials:
        m = await _db.inv_smart_meters.find_one({"serial_number": sn,
                                                   "deleted_at": {"$exists": False}})
        if not m:
            invalid.append(f"{sn}: not in stock")
        elif m["status"] in ("Issued", "Installed", "Damaged", "Returned"):
            invalid.append(f"{sn}: {m['status']}")
    if invalid:
        raise HTTPException(400, "Serial validation failed: " + "; ".join(invalid))

    doc = {"id": str(uuid.uuid4()), **payload.model_dump(),
           "meter_serials": serials,
           "created_at": _now(), "created_by": user["email"]}
    await _db.inv_gate_passes.insert_one(doc)

    # Auto-movement: on Approved, mark meters Issued
    if payload.status == "Approved" and serials:
        await _db.inv_smart_meters.update_many(
            {"serial_number": {"$in": serials}},
            {"$set": {"status": "Issued", "gate_pass_id": doc["id"], "issued_at": _now()}}
        )
        for sn in serials:
            m = await _db.inv_smart_meters.find_one({"serial_number": sn}, {"id": 1})
            if m:
                await _ledger(user["email"], "smart_meter", m["id"], "issued",
                               None, {"status": "Issued", "gate_pass": payload.gate_pass_number})

    await _ledger(user["email"], "gate_pass", doc["id"], "created", None,
                   {"number": payload.gate_pass_number, "meters": len(serials)})
    return _clean(doc)


@inventory_router.patch("/gate-passes/{gid}")
async def update_gate_pass(gid: str, patch: Dict[str, Any], user: dict = Depends(dep_admin)):
    before = await _db.inv_gate_passes.find_one({"id": gid}, {"_id": 0})
    if not before: raise HTTPException(404, "Not found")
    updates = {k: v for k, v in patch.items() if k not in ("id", "_id", "created_at", "created_by")}
    updates["updated_at"] = _now()
    await _db.inv_gate_passes.update_one({"id": gid}, {"$set": updates})
    after = await _db.inv_gate_passes.find_one({"id": gid}, {"_id": 0})
    # If status transitioned to Approved, issue the meters
    if before.get("status") != "Approved" and after.get("status") == "Approved":
        serials = after.get("meter_serials", [])
        if serials:
            await _db.inv_smart_meters.update_many(
                {"serial_number": {"$in": serials}, "status": "Available"},
                {"$set": {"status": "Issued", "gate_pass_id": gid, "issued_at": _now()}})
            for sn in serials:
                m = await _db.inv_smart_meters.find_one({"serial_number": sn}, {"id": 1})
                if m:
                    await _ledger(user["email"], "smart_meter", m["id"], "issued",
                                   None, {"status": "Issued"}, f"GP {after['gate_pass_number']}")
    await _ledger(user["email"], "gate_pass", gid, "updated", before, after)
    return after


@inventory_router.get("/gate-passes")
async def list_gate_passes(
    q: Optional[str] = None, status: Optional[str] = None,
    page: int = 1, page_size: int = Query(25, le=200),
    user: dict = Depends(dep_current),
):
    query = {"deleted_at": {"$exists": False}}
    if status: query["status"] = status
    if q:
        rex = {"$regex": q, "$options": "i"}
        query["$or"] = [{"gate_pass_number": rex}, {"vehicle_number": rex},
                        {"agency": rex}, {"to_location": rex}]
    skip = (max(1, page) - 1) * page_size
    total = await _db.inv_gate_passes.count_documents(query)
    docs = await _db.inv_gate_passes.find(query, {"_id": 0}).sort("gate_pass_date", -1).skip(skip).limit(page_size).to_list(page_size)
    return {"items": docs, "total": total, "page": page, "page_size": page_size}


@inventory_router.get("/gate-passes/{gid}")
async def get_gate_pass(gid: str, user: dict = Depends(dep_current)):
    doc = await _db.inv_gate_passes.find_one({"id": gid}, {"_id": 0})
    if not doc: raise HTTPException(404, "Not found")
    return doc


# ---------- CABLE ISSUE ----------
class CableIssueIn(BaseModel):
    issue_date: str
    division: Optional[str] = None
    sub_division: Optional[str] = None
    sdo: Optional[str] = None
    cable_type: str
    cable_size: str
    drum_number: str
    quantity: float
    issued_to: Optional[str] = None
    agency: Optional[str] = None
    work_order: Optional[str] = None
    vehicle_number: Optional[str] = None
    issue_slip_number: Optional[str] = None
    document: Optional[str] = None
    document_name: Optional[str] = None
    remarks: Optional[str] = None


@inventory_router.post("/cable-issues")
async def add_cable_issue(payload: CableIssueIn, user: dict = Depends(dep_admin)):
    drum = payload.drum_number.strip().upper()
    cable = await _db.inv_cables.find_one({"drum_number": drum, "deleted_at": {"$exists": False}})
    if not cable:
        raise HTTPException(400, f"Drum {drum} not in stock")
    balance = _cable_balance(cable)
    if payload.quantity <= 0:
        raise HTTPException(400, "Quantity must be positive")
    if payload.quantity > balance + 0.001:
        raise HTTPException(400, f"Only {balance} available on drum {drum}")

    doc = {"id": str(uuid.uuid4()), **payload.model_dump(),
           "drum_number": drum, "cable_id": cable["id"],
           "created_at": _now(), "created_by": user["email"]}
    await _db.inv_cable_issues.insert_one(doc)

    # Auto-movement: increment issued_qty on cable
    new_issued = (cable.get("issued_qty") or 0) + payload.quantity
    new_bal = _cable_balance({**cable, "issued_qty": new_issued})
    await _db.inv_cables.update_one({"id": cable["id"]},
        {"$set": {"issued_qty": new_issued, "balance_qty": new_bal}})
    await _ledger(user["email"], "cable", cable["id"], "issued",
                   {"balance": balance}, {"balance": new_bal},
                   f"Issue {payload.quantity} {cable.get('unit','')} to {payload.agency or payload.issued_to}")
    await _ledger(user["email"], "cable_issue", doc["id"], "created", None,
                   {"drum": drum, "qty": payload.quantity})
    return _clean(doc)


@inventory_router.get("/cable-issues")
async def list_cable_issues(
    q: Optional[str] = None, page: int = 1, page_size: int = Query(25, le=200),
    user: dict = Depends(dep_current),
):
    query = {"deleted_at": {"$exists": False}}
    if q:
        rex = {"$regex": q, "$options": "i"}
        query["$or"] = [{"drum_number": rex}, {"issue_slip_number": rex},
                        {"agency": rex}, {"work_order": rex}]
    skip = (max(1, page) - 1) * page_size
    total = await _db.inv_cable_issues.count_documents(query)
    docs = await _db.inv_cable_issues.find(query, {"_id": 0}).sort("issue_date", -1).skip(skip).limit(page_size).to_list(page_size)
    return {"items": docs, "total": total, "page": page, "page_size": page_size}


# ---------- BISIGNOFF (Used Cable) ----------
class BISignoffIn(BaseModel):
    bisignoff_number: str
    bisignoff_date: str
    division: Optional[str] = None
    sub_division: Optional[str] = None
    sdo: Optional[str] = None
    consumer_reference: Optional[str] = None
    cable_type: str
    cable_size: str
    drum_number: Optional[str] = None
    issued_qty: float = 0
    used_qty: float = 0
    balance_return: float = 0
    installer: Optional[str] = None
    agency: Optional[str] = None
    work_location: Optional[str] = None
    installation_id: Optional[str] = None
    status: Literal["Pending", "Submitted", "Verified", "Rejected"] = "Submitted"
    document: Optional[str] = None
    document_name: Optional[str] = None
    remarks: Optional[str] = None


@inventory_router.post("/bisignoffs")
async def add_bisignoff(payload: BISignoffIn, user: dict = Depends(dep_admin)):
    if await _db.inv_bisignoffs.find_one({"bisignoff_number": payload.bisignoff_number.strip(),
                                            "deleted_at": {"$exists": False}}):
        raise HTTPException(409, "Duplicate BISignoff number")
    doc = {"id": str(uuid.uuid4()), **payload.model_dump(),
           "created_at": _now(), "created_by": user["email"]}
    await _db.inv_bisignoffs.insert_one(doc)
    await _ledger(user["email"], "bisignoff", doc["id"], "created", None,
                   {"number": payload.bisignoff_number})
    return _clean(doc)


@inventory_router.patch("/bisignoffs/{bid}")
async def update_bisignoff(bid: str, patch: Dict[str, Any], user: dict = Depends(dep_admin)):
    before = await _db.inv_bisignoffs.find_one({"id": bid}, {"_id": 0})
    if not before: raise HTTPException(404, "Not found")
    # Verified is locked
    if before.get("status") == "Verified" and patch.get("status") != "Verified":
        raise HTTPException(400, "Verified BISignoff cannot be modified")
    if before.get("status") == "Verified":
        allowed_after_verify = {"remarks"}
        patch = {k: v for k, v in patch.items() if k in allowed_after_verify}
        if not patch:
            raise HTTPException(400, "Verified BISignoff is locked")
    updates = {k: v for k, v in patch.items() if k not in ("id", "_id", "created_at", "created_by")}
    if patch.get("status") == "Verified":
        updates["verified_by"] = user["email"]
        updates["verified_at"] = _now()
    updates["updated_at"] = _now()
    await _db.inv_bisignoffs.update_one({"id": bid}, {"$set": updates})
    after = await _db.inv_bisignoffs.find_one({"id": bid}, {"_id": 0})

    # Auto-movement: on Verified, update used_qty on cable drum
    if before.get("status") != "Verified" and after.get("status") == "Verified" and after.get("drum_number"):
        drum = after["drum_number"].strip().upper()
        cable = await _db.inv_cables.find_one({"drum_number": drum})
        if cable and after.get("used_qty"):
            new_used = (cable.get("used_qty") or 0) + float(after["used_qty"])
            await _db.inv_cables.update_one({"id": cable["id"]},
                {"$set": {"used_qty": new_used,
                          "balance_qty": _cable_balance({**cable, "used_qty": new_used})}})
            await _ledger(user["email"], "cable", cable["id"], "used",
                           None, {"used_qty": new_used},
                           f"BISignoff {after['bisignoff_number']} verified")

    await _ledger(user["email"], "bisignoff", bid, "updated", before, after)
    return after


@inventory_router.get("/bisignoffs")
async def list_bisignoffs(
    q: Optional[str] = None, status: Optional[str] = None,
    page: int = 1, page_size: int = Query(25, le=200),
    user: dict = Depends(dep_current),
):
    query = {"deleted_at": {"$exists": False}}
    if status: query["status"] = status
    if q:
        rex = {"$regex": q, "$options": "i"}
        query["$or"] = [{"bisignoff_number": rex}, {"drum_number": rex},
                        {"consumer_reference": rex}, {"work_location": rex}]
    skip = (max(1, page) - 1) * page_size
    total = await _db.inv_bisignoffs.count_documents(query)
    docs = await _db.inv_bisignoffs.find(query, {"_id": 0}).sort("bisignoff_date", -1).skip(skip).limit(page_size).to_list(page_size)
    return {"items": docs, "total": total, "page": page, "page_size": page_size}


# ---------- CROSS-LINKED HISTORY (mini Session C) ----------
@inventory_router.get("/history/serial/{sn}")
async def cross_history(sn: str, user: dict = Depends(dep_current)):
    """Return the FULL cross-linked lifecycle of a smart meter serial number."""
    sn = sn.strip().upper()
    meter = await _db.inv_smart_meters.find_one({"serial_number": sn}, {"_id": 0})
    if not meter: raise HTTPException(404, "Serial not found in stock")

    installation = await _db.inv_installations.find_one({"new_meter_serial": sn}, {"_id": 0})
    old_meter = None
    if installation and installation.get("old_meter_serial"):
        old_meter = await _db.inv_old_meters.find_one({"serial_number": installation["old_meter_serial"]}, {"_id": 0})

    gate_pass = None
    if meter.get("gate_pass_id"):
        gate_pass = await _db.inv_gate_passes.find_one({"id": meter["gate_pass_id"]}, {"_id": 0})

    bisignoff = None
    if installation:
        bisignoff = await _db.inv_bisignoffs.find_one({"installation_id": installation["id"]}, {"_id": 0})

    ledger = await _db.inv_ledger.find({"entity": "smart_meter", "entity_id": meter["id"]},
                                         {"_id": 0}).sort("timestamp", 1).to_list(200)

    return {
        "serial": sn, "meter": meter, "gate_pass": gate_pass,
        "installation": installation, "old_meter": old_meter,
        "bisignoff": bisignoff, "ledger": ledger,
    }


async def ensure_inventory_indexes(db):
    await db.inv_master.create_index([("type", 1), ("name", 1)], unique=False)
    await db.inv_smart_meters.create_index("serial_number", unique=False)
    await db.inv_smart_meters.create_index("status")
    await db.inv_smart_meters.create_index("division")
    await db.inv_old_meters.create_index("serial_number")
    await db.inv_old_meters.create_index("deposit_status")
    await db.inv_cables.create_index("drum_number")
    await db.inv_ledger.create_index([("timestamp", -1)])
    await db.inv_ledger.create_index([("entity", 1), ("entity_id", 1)])
    await db.inv_installations.create_index("new_meter_serial")
    await db.inv_installations.create_index("consumer_number")
    await db.inv_gate_passes.create_index("gate_pass_number")
    await db.inv_gate_passes.create_index("status")
    await db.inv_cable_issues.create_index("drum_number")
    await db.inv_bisignoffs.create_index("bisignoff_number")
    await db.inv_bisignoffs.create_index("status")


# ================================================================
#            SESSION C — EXCEL EXPORTS (per collection)
# ================================================================
from fastapi.responses import StreamingResponse

REPORT_MAP = {
    "smart-meters": ("inv_smart_meters", "Smart Meters",
                      ["serial_number", "meter_make", "meter_model", "meter_type", "rating",
                       "batch_number", "purchase_date", "received_qty", "status",
                       "division", "sub_division", "sdo", "store_location", "vendor",
                       "gate_pass_id", "installation_id", "created_at", "created_by", "remarks"]),
    "old-meters": ("inv_old_meters", "Old Meters",
                    ["serial_number", "consumer_number", "consumer_name", "meter_make",
                     "meter_model", "removal_date", "condition", "deposit_status",
                     "deposit_date", "division", "sub_division", "sdo", "installer",
                     "store_location", "created_at", "remarks"]),
    "cables": ("inv_cables", "Cables",
                ["drum_number", "cable_type", "cable_size", "cable_spec", "make",
                 "batch_number", "unit", "opening_stock", "received_qty", "issued_qty",
                 "used_qty", "returned_qty", "damaged_qty", "balance_qty",
                 "division", "sub_division", "sdo", "store_location", "receipt_date", "remarks"]),
    "installations": ("inv_installations", "Installations",
                       ["installation_date", "consumer_name", "consumer_number", "address",
                        "new_meter_serial", "old_meter_serial", "division", "sub_division",
                        "sdo", "installer", "agency", "mobile_number", "status",
                        "latitude", "longitude", "created_by", "remarks"]),
    "gate-passes": ("inv_gate_passes", "Gate Passes",
                     ["gate_pass_number", "gate_pass_date", "from_location", "to_location",
                      "vehicle_number", "driver_name", "driver_mobile", "agency", "purpose",
                      "prepared_by", "approved_by", "status", "meter_count", "created_by"]),
    "cable-issues": ("inv_cable_issues", "Cable Issues",
                      ["issue_date", "issue_slip_number", "drum_number", "cable_type",
                       "cable_size", "quantity", "issued_to", "agency", "work_order",
                       "vehicle_number", "created_by", "remarks"]),
    "bisignoffs": ("inv_bisignoffs", "BISignoffs",
                    ["bisignoff_number", "bisignoff_date", "consumer_reference",
                     "cable_type", "cable_size", "drum_number", "issued_qty", "used_qty",
                     "balance_return", "installer", "agency", "work_location", "status",
                     "verified_by", "verified_at", "created_by", "remarks"]),
    "ledger": ("inv_ledger", "Audit Trail",
                ["timestamp", "entity", "entity_id", "action", "actor", "detail"]),
}


@inventory_router.get("/reports/{kind}.xlsx")
async def export_xlsx(kind: str, user: dict = Depends(dep_current),
                       start: Optional[str] = None, end: Optional[str] = None,
                       division: Optional[str] = None):
    if kind not in REPORT_MAP:
        raise HTTPException(400, f"Unknown report kind. Options: {sorted(REPORT_MAP)}")
    coll_name, sheet_name, cols = REPORT_MAP[kind]
    coll = getattr(_db, coll_name)
    query = {} if coll_name == "inv_ledger" else {"deleted_at": {"$exists": False}}
    # Determine the date field per collection
    date_field = {
        "inv_ledger": "timestamp",
        "inv_installations": "installation_date",
        "inv_gate_passes": "gate_pass_date",
        "inv_cable_issues": "issue_date",
        "inv_bisignoffs": "bisignoff_date",
    }.get(coll_name, "created_at")
    if start or end:
        rq = {}
        if start: rq["$gte"] = start
        if end: rq["$lte"] = end + "T23:59:59"
        query[date_field] = rq
    if division:
        query["division"] = division
    sort_field = "timestamp" if coll_name == "inv_ledger" else "created_at"
    docs = await coll.find(query, {"_id": 0}).sort(sort_field, -1).to_list(20000)

    from openpyxl.styles import Font, PatternFill, Alignment
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_name[:31]
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="0B1E3F", end_color="0B1E3F", fill_type="solid")
    for i, c in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i, value=c.replace("_", " ").upper())
        cell.font = header_font; cell.fill = header_fill
        cell.alignment = Alignment(horizontal="left", vertical="center")
    for r_idx, d in enumerate(docs, start=2):
        for c_idx, c in enumerate(cols, start=1):
            v = d.get(c)
            if c == "meter_count":
                v = len(d.get("meter_serials") or [])
            if isinstance(v, (dict, list)):
                v = str(v)[:200]
            ws.cell(row=r_idx, column=c_idx, value=v)
    for col_cells in ws.columns:
        length = max((len(str(cell.value)) if cell.value is not None else 0) for cell in col_cells)
        ws.column_dimensions[col_cells[0].column_letter].width = min(max(length + 2, 12), 40)
    ws.freeze_panes = "A2"

    buf = io.BytesIO(); wb.save(buf); buf.seek(0)
    filename = f"pps-inventory-{kind}-{datetime.now(timezone.utc).strftime('%Y%m%d-%H%M')}.xlsx"
    return StreamingResponse(buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@inventory_router.get("/reports/summary")
async def report_summary(user: dict = Depends(dep_current)):
    out = []
    for kind, (coll_name, sheet_name, _) in REPORT_MAP.items():
        coll = getattr(_db, coll_name)
        query = {} if coll_name == "inv_ledger" else {"deleted_at": {"$exists": False}}
        total = await coll.count_documents(query)
        sort_field = "timestamp" if coll_name == "inv_ledger" else "created_at"
        latest = await coll.find(query, {"_id": 0}).sort(sort_field, -1).limit(1).to_list(1)
        latest_ts = None
        if latest:
            latest_ts = latest[0].get("timestamp") or latest[0].get("created_at")
        out.append({"kind": kind, "name": sheet_name, "rows": total, "last_updated": latest_ts})
    return {"reports": out}


# ================================================================
#            BRANDED HTML PRINT (Gate Pass / BISignoff)
# ================================================================
from fastapi.responses import HTMLResponse


def _print_header_html():
    return """
    <style>
      * { box-sizing: border-box; margin: 0; padding: 0; }
      body { font-family: 'Arial', 'Helvetica', sans-serif; color: #111; background: #fff; padding: 24px; font-size: 12px; line-height: 1.4; }
      .header { border-bottom: 3px solid #0B1E3F; padding-bottom: 12px; margin-bottom: 16px; display: flex; justify-content: space-between; align-items: flex-start; }
      .brand { display: flex; align-items: center; gap: 12px; }
      .logo { width: 48px; height: 48px; background: #0B1E3F; color: #fff; border-radius: 8px; display: flex; align-items: center; justify-content: center; font-weight: 900; font-size: 22px; }
      .brand-text .name { font-size: 18px; font-weight: 900; color: #0B1E3F; letter-spacing: 0.5px; }
      .brand-text .tag { font-size: 10px; letter-spacing: 2px; color: #94A3B8; text-transform: uppercase; font-weight: 700; margin-top: 2px; }
      .doc-title { text-align: right; }
      .doc-title h1 { font-size: 22px; font-weight: 900; color: #0B1E3F; letter-spacing: 1px; }
      .doc-title .num { font-family: 'Courier New', monospace; font-size: 13px; color: #444; margin-top: 4px; }
      .grid { display: grid; grid-template-columns: 1fr 1fr; gap: 8px 24px; margin: 14px 0; }
      .field { border-bottom: 1px dotted #e5e7eb; padding: 6px 0; }
      .field .k { font-size: 9px; letter-spacing: 1.2px; text-transform: uppercase; color: #64748b; font-weight: 700; }
      .field .v { font-size: 12px; font-weight: 600; margin-top: 2px; }
      table.tbl { width: 100%; border-collapse: collapse; margin: 14px 0; }
      table.tbl th { background: #0B1E3F; color: #fff; font-size: 10px; letter-spacing: 1px; padding: 8px; text-align: left; text-transform: uppercase; }
      table.tbl td { padding: 6px 8px; border-bottom: 1px solid #e5e7eb; font-size: 11px; }
      table.tbl tr:nth-child(even) td { background: #f8fafc; }
      .footer { margin-top: 32px; padding-top: 16px; border-top: 1px solid #e5e7eb; display: grid; grid-template-columns: 1fr 1fr 1fr; gap: 32px; font-size: 10px; }
      .sig-box { border-top: 1px solid #111; padding-top: 6px; font-weight: 700; text-transform: uppercase; letter-spacing: 1px; margin-top: 48px; text-align: center; }
      .status-badge { display: inline-block; padding: 4px 12px; border-radius: 999px; font-size: 10px; font-weight: 900; letter-spacing: 1px; text-transform: uppercase; }
      .stamp { color: #22C55E; }
      .meta { color: #64748b; font-size: 10px; margin-top: 20px; text-align: center; }
      @media print { body { padding: 12mm; } .no-print { display: none; } }
    </style>
    """


def _print_brand_block():
    return """
    <div class="brand">
      <div class="logo">P</div>
      <div class="brand-text">
        <div class="name">PRATHVI POWER SOLUTIONS</div>
        <div class="tag">Reliable Power · Smarter Solutions</div>
      </div>
    </div>
    """


def _field(k, v):
    v = v if v not in (None, "", []) else "—"
    return f'<div class="field"><div class="k">{k}</div><div class="v">{v}</div></div>'


@inventory_router.get("/gate-passes/{gid}/print", response_class=HTMLResponse)
async def print_gate_pass(gid: str, user: dict = Depends(dep_current)):
    doc = await _db.inv_gate_passes.find_one({"id": gid}, {"_id": 0})
    if not doc: raise HTTPException(404, "Not found")
    serials = doc.get("meter_serials") or []
    serial_rows = "".join(
        f"<tr><td>{i+1}</td><td style='font-family:monospace'>{s}</td></tr>"
        for i, s in enumerate(serials)
    ) or "<tr><td colspan='2' style='text-align:center;color:#94a3b8'>No serials listed</td></tr>"

    html = f"""<!DOCTYPE html><html><head><meta charset="utf-8"><title>Gate Pass {doc.get('gate_pass_number')}</title>
{_print_header_html()}</head><body>
<div class="header">
  {_print_brand_block()}
  <div class="doc-title">
    <h1>GATE PASS</h1>
    <div class="num">No. {doc.get('gate_pass_number','—')}</div>
    <div class="num">Date: {doc.get('gate_pass_date','—')}</div>
  </div>
</div>
<div class="grid">
  {_field('From Location', doc.get('from_location'))}
  {_field('To Location', doc.get('to_location'))}
  {_field('Division', doc.get('division'))}
  {_field('Sub Division', doc.get('sub_division'))}
  {_field('SDO', doc.get('sdo'))}
  {_field('Agency / Vendor', doc.get('agency'))}
  {_field('Vehicle Number', doc.get('vehicle_number'))}
  {_field('Driver Name', doc.get('driver_name'))}
  {_field('Driver Mobile', doc.get('driver_mobile'))}
  {_field('Purpose', doc.get('purpose'))}
  {_field('Prepared By', doc.get('prepared_by'))}
  {_field('Approved By', doc.get('approved_by'))}
</div>
<div style="margin-top:8px"><span class="status-badge" style="background:#0EA5E9;color:#fff">Status · {doc.get('status','—').upper()}</span></div>
<h3 style="margin-top:20px;font-size:13px;color:#0B1E3F;letter-spacing:1px">METER SERIALS ({len(serials)})</h3>
<table class="tbl"><thead><tr><th style="width:60px">#</th><th>Serial Number</th></tr></thead><tbody>{serial_rows}</tbody></table>
{f'<div>Remarks: <em>{doc.get("remarks")}</em></div>' if doc.get('remarks') else ''}
<div class="footer">
  <div class="sig-box">Prepared By</div>
  <div class="sig-box">Store In-Charge</div>
  <div class="sig-box">Received By</div>
</div>
<div class="meta">Generated on {datetime.now(timezone.utc).astimezone().strftime('%d %b %Y · %I:%M %p')} · PPS Inventory · This is a system-generated document.</div>
<script>window.onload=()=>setTimeout(()=>window.print(),300);</script>
</body></html>"""
    return HTMLResponse(content=html)


@inventory_router.get("/bisignoffs/{bid}/print", response_class=HTMLResponse)
async def print_bisignoff(bid: str, user: dict = Depends(dep_current)):
    doc = await _db.inv_bisignoffs.find_one({"id": bid}, {"_id": 0})
    if not doc: raise HTTPException(404, "Not found")
    status_color = "#22C55E" if doc.get("status") == "Verified" else "#F59E0B"
    verified_stamp = ""
    if doc.get("status") == "Verified":
        verified_stamp = f"""<div style="position:absolute;top:110px;right:60px;border:4px solid #22C55E;color:#22C55E;padding:8px 24px;font-weight:900;font-size:16px;letter-spacing:3px;transform:rotate(-8deg);opacity:0.85;border-radius:4px">VERIFIED</div>"""

    html = f"""<!DOCTYPE html><html><head><meta charset="utf-8"><title>BISignoff {doc.get('bisignoff_number')}</title>
{_print_header_html()}</head><body style="position:relative">
{verified_stamp}
<div class="header">
  {_print_brand_block()}
  <div class="doc-title">
    <h1>BI-SIGNOFF</h1>
    <div class="num">No. {doc.get('bisignoff_number','—')}</div>
    <div class="num">Date: {doc.get('bisignoff_date','—')}</div>
  </div>
</div>
<div class="grid">
  {_field('Consumer / Work Reference', doc.get('consumer_reference'))}
  {_field('Work Location', doc.get('work_location'))}
  {_field('Division', doc.get('division'))}
  {_field('Sub Division', doc.get('sub_division'))}
  {_field('SDO', doc.get('sdo'))}
  {_field('Installer', doc.get('installer'))}
  {_field('Agency', doc.get('agency'))}
</div>
<h3 style="margin-top:20px;font-size:13px;color:#0B1E3F;letter-spacing:1px">CABLE CONSUMPTION</h3>
<table class="tbl">
  <thead><tr><th>Cable Type</th><th>Size</th><th>Drum No</th><th>Issued</th><th>Used</th><th>Balance / Return</th></tr></thead>
  <tbody><tr>
    <td>{doc.get('cable_type','—')}</td><td>{doc.get('cable_size','—')}</td>
    <td style="font-family:monospace">{doc.get('drum_number','—')}</td>
    <td>{doc.get('issued_qty',0)}</td>
    <td><strong>{doc.get('used_qty',0)}</strong></td>
    <td>{doc.get('balance_return',0)}</td>
  </tr></tbody>
</table>
<div style="margin-top:8px"><span class="status-badge" style="background:{status_color};color:#fff">Status · {doc.get('status','—').upper()}</span>
{f'<span style="margin-left:12px;font-size:10px;color:#64748b">Verified by {doc.get("verified_by")} on {doc.get("verified_at","")[:10]}</span>' if doc.get('status') == 'Verified' else ''}
</div>
{f'<div style="margin-top:14px">Remarks: <em>{doc.get("remarks")}</em></div>' if doc.get('remarks') else ''}
<div class="footer">
  <div class="sig-box">Installer</div>
  <div class="sig-box">Consumer / Owner</div>
  <div class="sig-box">Verifier</div>
</div>
<div class="meta">Generated on {datetime.now(timezone.utc).astimezone().strftime('%d %b %Y · %I:%M %p')} · PPS Inventory · This is a system-generated document.</div>
<script>window.onload=()=>setTimeout(()=>window.print(),300);</script>
</body></html>"""
    return HTMLResponse(content=html)


async def seed_inventory_masters(db):
    """Seed a minimal set of common masters if empty."""
    count = await db.inv_master.count_documents({})
    if count > 0:
        return
    seeds = [
        ("meter_make", ["L&T", "HPL", "Genus", "Secure", "Schneider", "Iskraemeco"]),
        ("meter_model", ["Single Phase Smart", "Three Phase Smart", "LT-CT Smart"]),
        ("cable_type", ["Aluminium", "Copper", "ABC"]),
        ("cable_size", ["16 sq.mm", "25 sq.mm", "50 sq.mm", "95 sq.mm", "185 sq.mm"]),
        ("division", ["Sitapur Division-I", "Sitapur Division-II", "Biswan Division-III", "Mahmudabad Division-IV"]),
        ("document_type", ["Gate Pass", "BISignoff", "Issue Slip", "Return Slip", "Photo"]),
    ]
    docs = []
    for t, names in seeds:
        for n in names:
            docs.append({"id": str(uuid.uuid4()), "type": t, "name": n,
                          "active": True, "created_at": _now(), "created_by": "seed"})
    if docs:
        await db.inv_master.insert_many(docs)
        log.info(f"Seeded {len(docs)} inventory master rows")
